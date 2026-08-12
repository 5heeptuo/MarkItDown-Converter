"""转换核心：MarkItDown 优先，空内容时本地 OCR fallback，错误分类。

流程：
    convert_file(input, output, stage_cb)
      ├─ 格式检查（markitdown / ocr-image / unknown）
      ├─ MarkItDown 转换
      ├─ 空内容检测
      │    ├─ 图片 -> OCR（jpg/jpeg/png/gif/bmp/webp/tiff/svg）
      │    ├─ PDF  -> 扫描版 OCR（渲染每页）
      │    └─ 其他 -> NO_CONTENT
      └─ 保存 Markdown
"""

from __future__ import annotations

import enum
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Callable

from format_registry import OCR_IMAGE_EXTENSIONS, classify
from logger import log_exception


class Status(str, enum.Enum):
    """转换状态枚举。"""

    WAITING = "等待"
    CONVERTING = "转换中"
    OCR = "OCR 识别中"
    SUCCESS = "完成"
    UNSUPPORTED = "不支持"
    NO_CONTENT = "无内容"
    FAILED = "失败"
    SKIPPED = "跳过"

    def __str__(self) -> str:  # noqa: D105
        return self.value


class ConvertError(Exception):
    """带友好消息的转换错误。message 可直接展示给用户（不含 traceback）。"""

    def __init__(self, message: str, category: str) -> None:
        super().__init__(message)
        self.category = category


# MarkItDown 单例（懒加载；子进程各自加载，互不影响）
_markitdown_instance = None


def _get_markitdown():
    global _markitdown_instance
    if _markitdown_instance is None:
        from markitdown import MarkItDown

        _markitdown_instance = MarkItDown()
    return _markitdown_instance


def _classify_error(exc: BaseException, src: Path, dst: Path) -> ConvertError:
    """把异常分类为用户可读的错误消息。完整 traceback 由调用方写日志。"""
    if _is_corrupt_chain(exc):
        return ConvertError("文件可能已损坏或格式不正确", "corrupt")
    if isinstance(exc, FileNotFoundError):
        return ConvertError("无法读取文件：文件不存在", "read")
    if isinstance(exc, PermissionError):
        return ConvertError("没有读取权限，请检查文件或文件夹权限", "permission")
    if isinstance(exc, IsADirectoryError):
        return ConvertError("无法读取文件：这是一个文件夹而不是文件", "read")
    # 输出目录写入失败
    if not dst.parent.exists() or not _dir_writable(dst.parent):
        return ConvertError("无法写入输出目录，请检查目录权限", "output")
    return ConvertError("MarkItDown 转换失败", "convert")


def _is_corrupt_chain(exc: BaseException) -> bool:
    """沿异常链查找损坏类异常（markitdown 会包装底层异常）。"""
    message = f"{type(exc).__name__} {exc}".lower()
    if any(token in message for token in (
        "badzipfile", "is not a zip file", "not a zip",
        "parseerror", "failed to parse", "unmarshal error",
    )):
        return True
    seen: set[int] = set()
    current: BaseException | None = exc
    while current is not None and id(current) not in seen:
        seen.add(id(current))
        if isinstance(current, (zipfile.BadZipFile, ET.ParseError, UnicodeDecodeError)):
            return True
        if "BadZipFile" in type(current).__name__ or "ParseError" in type(current).__name__:
            return True
        current = current.__cause__ or current.__context__
    return False


def _dir_writable(directory: Path) -> bool:
    try:
        probe = directory / ".mdc_write_probe"
        probe.write_text("", encoding="utf-8")
        probe.unlink(missing_ok=True)
        return True
    except OSError:
        return False


def _is_empty(text: str | None) -> bool:
    return not text or not text.strip()


def _xlsx_has_real_nan_text(src: Path) -> bool:
    """源 XLSX 中是否真实存在文本 'NaN'/'nan'/'None' 单元格。

    存在时保守地不做替换（避免误删用户真正写下的文本）。
    """
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.strip() in (
                            "NaN",
                            "nan",
                            "None",
                        ):
                            return True
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        return True  # 读不到源文件时保守不替换
    return False


_TABLE_NAN_CELL = re.compile(r"(?<=[|\n])\s*(?:NaN|nan|None)\s*(?=[|\n])")


def _clean_table_nan(text: str) -> str:
    """把 Markdown 表格单元格中孤立的 NaN/nan/None 替换为空（不碰正文）。"""
    return _TABLE_NAN_CELL.sub("", text)


def _sanitize_zip_path(text: str, zip_name: str) -> str:
    """ZIP Markdown 不暴露本机绝对路径：只保留归档文件名。"""
    lines = text.splitlines()
    for i, line in enumerate(lines):
        if line.startswith("Content from the zip file"):
            lines[i] = f"Content from the zip file `{zip_name}`:"
    return "\n".join(lines)


def _ocr_image(src: Path) -> str:
    """图片 OCR（含 svg 栅格化尝试）。"""
    if src.suffix.lower() == ".svg":
        text = _ocr_svg(src)
        if text:
            return text
        return ""
    from ocr import ocr_image

    return ocr_image(src)


def _ocr_svg(src: Path) -> str:
    """SVG 优先提取内嵌 <text> 文本；否则尝试栅格化后 OCR。"""
    text = _extract_svg_text(src)
    if text:
        return text
    try:
        import cairosvg

        from PIL import Image

        from ocr import ocr_image_from_pil

        png_path = src.with_suffix(".ocr_probe.png")
        cairosvg.svg2png(url=str(src), write_to=str(png_path), output_width=1600)
        try:
            img = Image.open(png_path)
            return ocr_image_from_pil(img)
        finally:
            png_path.unlink(missing_ok=True)
    except Exception as exc:  # noqa: BLE001
        log_exception(exc, f"SVG 栅格化失败: {src}")
        return ""


def _extract_svg_text(src: Path) -> str:
    """从 SVG XML 中提取 <text> 元素内容（SVG 本质是文本格式，多数含文字）。"""
    try:
        root = ET.parse(src).getroot()
        parts: list[str] = []
        for elem in root.iter():
            tag = elem.tag.rsplit("}", 1)[-1].lower()
            if tag == "text" and elem.text and elem.text.strip():
                parts.append(elem.text.strip())
        return "\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        log_exception(exc, f"SVG 文本提取失败: {src}")
        return ""


def convert_file(
    input_path: str | Path,
    output_path: str | Path,
    stage_cb: Callable[[Status], None] | None = None,
) -> tuple[Status, str]:
    """统一转换入口。

    Args:
        input_path:  输入文件
        output_path: 输出 .md 文件
        stage_cb:    阶段回调（例如 GUI 显示“OCR 识别中”）

    Returns:
        (Status, message)。message 为空表示无附加信息。
    """
    src = Path(input_path)
    dst = Path(output_path)

    # ------------------------------------------------------------- 前置检查
    if not src.exists():
        raise ConvertError("无法读取文件：文件不存在", "read")
    if not src.is_file():
        raise ConvertError("无法读取文件：这是一个文件夹而不是文件", "read")
    if src.suffix.lower() == ".lnk":
        raise ConvertError("不支持快捷方式，请先解析 .lnk 文件", "unsupported")
    if classify(src) == "unsupported":
        # 明显不支持的格式（exe/dll/bat/cmd/ps1 等）：与 GUI 添加阶段一致，
        # 不交给 MarkItDown 按内容猜测（保证 CLI/GUI 行为一致）
        raise ConvertError(f"不支持的文件格式 .{src.suffix.lower().lstrip('.')}", "unsupported")

    # ------------------------------------------------------------- MarkItDown
    text = ""
    try:
        if stage_cb:
            stage_cb(Status.CONVERTING)
        result = _get_markitdown().convert(str(src))
        text = (result.text_content or "").strip()
        ext = src.suffix.lower()
        if ext == ".xlsx" and not _xlsx_has_real_nan_text(src):
            text = _clean_table_nan(text)
        elif ext == ".zip":
            text = _sanitize_zip_path(text, src.name)
    except Exception as exc:  # noqa: BLE001
        log_exception(exc, f"MarkItDown 转换失败: {src}")
        # MarkItDown 失败时，图片和 PDF 仍尝试 OCR fallback
        if src.suffix.lower() not in (".jpg", ".jpeg", ".png", ".pdf") \
                and src.suffix.lower() not in OCR_IMAGE_EXTENSIONS:
            raise _classify_error(exc, src, dst) from exc

    # ------------------------------------------------------------- OCR fallback
    if _is_empty(text):
        ext = src.suffix.lower()
        if ext in OCR_IMAGE_EXTENSIONS or ext in (".jpg", ".jpeg", ".png"):
            if stage_cb:
                stage_cb(Status.OCR)
            text = _ocr_image(src)
        elif ext == ".pdf":
            if stage_cb:
                stage_cb(Status.OCR)
            text = _ocr_pdf(src)
    elif src.suffix.lower() == ".pdf" and _pdf_text_suspicious(src, text):
        # 混合 PDF：MarkItDown 提取了少量文本（如每页不足一行），
        # 很可能存在无文本层的图片页 -> 用 OCR 补充
        if stage_cb:
            stage_cb(Status.OCR)
        ocr_text = _ocr_pdf(src)
        if not _is_empty(ocr_text):
            text = _merge_pdf_text(text, ocr_text)

    # ------------------------------------------------------------- 空内容检查
    if _is_empty(text):
        raise ConvertError("未检测到可提取内容（扫描件或图片可能没有可识别文字）", "no_content")

    # ------------------------------------------------------------- 保存
    try:
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_text(text, encoding="utf-8")
    except OSError as exc:  # noqa: BLE001
        log_exception(exc, f"写入输出失败: {dst}")
        raise ConvertError("无法写入输出目录，请检查目录权限", "output") from exc
    return Status.SUCCESS, ""


def _ocr_pdf(src: Path) -> str:
    from ocr import ocr_pdf

    return ocr_pdf(src)


def _pdf_text_suspicious(src: Path, text: str) -> bool:
    """判断 PDF 是否可能存在无文本层页面：文本行数不超过页数（平均每页不足一行）。"""
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return False  # 完全空的情况已由空内容分支处理
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(src))
        try:
            page_count = len(pdf)
        finally:
            pdf.close()
    except Exception:  # noqa: BLE001
        return False
    return page_count > 1 and len(lines) <= page_count


def _merge_pdf_text(primary: str, ocr_text: str) -> str:
    """合并 MarkItDown 文本与 OCR 文本：去重后拼接。"""
    primary_lines = [line for line in primary.splitlines() if line.strip()]
    primary_set = {line.strip() for line in primary_lines}
    extra = [line for line in ocr_text.splitlines() if line.strip() and line.strip() not in primary_set]
    if not extra:
        return primary
    return "\n\n".join([*primary_lines, *extra])


# 兼容旧接口：converter.convert_to_markdown(src, dst) -> None
def convert_to_markdown(input_path: str | Path, output_path: str | Path) -> None:
    """兼容包装：成功返回 None，失败抛 ConvertError。"""
    status, _ = convert_file(input_path, output_path)
    if status != Status.SUCCESS:
        raise ConvertError(f"转换未成功：{status.value}", status.name.lower())



